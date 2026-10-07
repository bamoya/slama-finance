"""Build the editable, client-facing business confirmation document."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Slama_Finance_Business_Rules_Approval.docx"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.7)
section.left_margin = section.right_margin = Cm(1.8)
section.footer_distance = Cm(0.8)

for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor.from_string("000000")
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.06
    style.paragraph_format.widow_control = True
doc.styles["Title"].font.size = Pt(25)
doc.styles["Title"].font.bold = True
doc.styles["Subtitle"].font.size = Pt(13)
doc.styles["Heading 1"].font.size = Pt(19)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(11)
doc.styles["Heading 2"].font.size = Pt(12)
doc.styles["Heading 2"].font.bold = True
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
    doc.styles[name].paragraph_format.keep_with_next = True
    pr = doc.styles[name].element.get_or_add_pPr()
    for border in list(pr.findall(qn("w:pBdr"))):
        pr.remove(border)

lang = OxmlElement("w:lang")
lang.set(qn("w:val"), "en-GB")
doc.styles["Normal"].element.get_or_add_rPr().append(lang)
doc.core_properties.title = "Slama Finance Business data and rules for client approval"
doc.core_properties.subject = "Business data and rules for client confirmation"
doc.core_properties.author = "Slama Finance"
doc.core_properties.keywords = "products, clients, estimates, invoices, delivery, payments"
doc.core_properties.comments = ""

# Page numbers make references in a returned approval form unambiguous.
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("Slama Finance  |  Business rules approval  |  ")
run.font.size = Pt(9)
run.font.color.rgb = RGBColor.from_string("666666")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para


def h(text):
    doc.add_heading(text, level=2)


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def page(title):
    doc.add_page_break()
    doc.add_heading(title, level=1)


def table(headers, rows, widths):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    for col, width in zip(tbl.columns, widths):
        col.width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    tbl._tbl.tblPr.append(borders)
    repeat = OxmlElement("w:tblHeader")
    tbl.rows[0]._tr.get_or_add_trPr().append(repeat)
    all_rows = [headers] + list(rows)
    for idx, values in enumerate(all_rows):
        row = tbl.rows[0] if idx == 0 else tbl.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, text in enumerate(values):
            cell = row.cells[j]
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "333333" if idx == 0 else ("F5F6F7" if idx % 2 == 0 else "FFFFFF"))
            tcpr.append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "start", "end"):
                n = OxmlElement(f"w:{side}")
                n.set(qn("w:w"), "85")
                n.set(qn("w:type"), "dxa")
                margins.append(n)
            tcpr.append(margins)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.03
            r = para.add_run(text)
            r.font.size = Pt(10.5)
            r.bold = idx == 0
            r.font.color.rgb = RGBColor.from_string("FFFFFF" if idx == 0 else "000000")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return tbl


def question(code, text):
    para = p(f"{code}  {text}", bold=True)
    para.paragraph_format.keep_with_next = True
    p("Response or clarification: ......................................................................................")


def confirm():
    p("Section:  ☐ Approved   ☐ Changes requested   ☐ Discuss", bold=True)


doc.add_paragraph("Slama Finance", style="Subtitle")
doc.add_paragraph("Business data and rules\nfor client approval", style="Title")
p("Version 1.0  •  26 September 2026")
p("Please confirm that the data and rules in this document match the way your company works. Your responses will establish the functional baseline for the first version of Slama Finance.")
p("Review each section, mark it as approved or requiring changes, and answer the confirmation questions. An unanswered question or unchecked box does not count as approval.")
h("What this document covers")
p("Products and categories, customer records, estimates, invoices, delivery notes and payment tracking. Staff management, notifications, reports and technical architecture are outside this review.")
h("Supported business flows")
table(["Starting point", "Possible result", "Main restriction"], [
    ("Direct entry", "An invoice", "No estimate is required beforehand."),
    ("One accepted estimate", "Several invoices over time", "Each invoice has at most one source estimate."),
    ("One invoice", "Several delivery notes", "Partial deliveries cannot exceed the invoiced weights."),
    ("Delivery before invoicing", "Combined or partial invoicing of delivered goods", "Same client; delivered weights cannot be billed twice."),
    ("One issued invoice", "Several payments", "Each payment belongs to one invoice only."),
], [4.2, 6.0, 7.2])
h("How to give your approval")
p("Use the section checkboxes and answer spaces. Refer to the question numbers when requesting changes. The final page provides room for your overall decision, comments, name and date.")
p("This document confirms software behavior. It does not change the commercial proposal or replace your accountant’s approval of invoice numbering and required document wording.")

page("1 Products and categories")
p("The catalog covers goods sold by weight. Products can be created, viewed, updated and archived when they should no longer be available for new sales.")
table(["Data", "Proposed rule"], [
    ("Reference and product name", "Required. Each product reference is unique."),
    ("Category", "Optional. One category per product. Categories can be created, renamed and archived."),
    ("Weight unit", "Grams (g) or kilograms (kg). Services, individual pieces, cartons and other sale units are not included."),
    ("Selling price", "Price excluding VAT per 1 g or 1 kg, according to the selected unit. Zero or positive."),
    ("Cost and description", "Optional. The description is internal catalog information."),
    ("Suggested VAT rate", "Optional. It never automatically applies VAT to a sale."),
], [4.3, 13.1])
h("Sales rules")
bullet("The actual weight is entered on each estimate, invoice or delivery line. For example, 2.5 kg at 12 DH/kg equals 30 DH excluding VAT.")
bullet("Products are searchable by name or reference. The PDF line displays the product name, not its internal description.")
bullet("Changing a product’s price, name or category does not update documents already issued. Archiving preserves historical records.")
bullet("Subcategories, multiple categories per product, inventory quantities and warehouse movements are not included.")
h("Please confirm")
question("P1", "Are g and kg sufficient, without sales by bag, carton or individual piece?")
question("P2", "Do you need discounts or different prices for different clients? These features are not included in the current rules.")
confirm()

page("2 Client data")
p("Two client types are supported: individual and company. A client does not need a user account or an email address for a sale to be recorded.")
table(["Data", "Individual", "Company"], [
    ("Billing identity", "First and last names required", "Legal company name required"),
    ("Trading name and contact", "Not included", "Optional"),
    ("Billing address", "Street address, city and country required", "Street address, city and country required"),
    ("Postal code", "Optional", "Optional"),
    ("Phone and email", "Optional", "Optional"),
    ("ICE, IF and professional tax ID", "Not included", "Can be added to the client record as available"),
    ("RC and registration city", "Not included", "Both provided together, or both left blank"),
    ("Delivery address", "Optional if the billing address is used", "Optional if the billing address is used"),
    ("Document language and notes", "French by default; notes optional", "French by default; notes optional"),
], [4.4, 6.5, 6.5])
h("Rules and restrictions")
bullet("Country defaults to Morocco. A different delivery address must contain its own street address, city and country.")
bullet("Updating a client record does not rewrite the client details on documents already issued. Archived clients remain available in historical records.")
bullet("CIN is not collected. Sole traders needing both personal and business identities require clarification before the two-type client model is approved.")
p("Fields needed to create a client record are not a declaration of legally required invoice information. The details required before issuing invoices will be confirmed with your accountant.")
h("Please confirm")
question("C1", "Do you invoice sole traders or walk-in customers without a complete address?")
question("C2", "Which identifiers and languages should appear on your documents? Please provide an existing example.")
confirm()

page("3 Estimates")
p("An estimate presents an offer to a client. It is separate from an invoice and does not record either a payment or a delivery.")
table(["Data", "Proposed rule"], [
    ("Client and dates", "Client, estimate date and validity deadline before issue."),
    ("Lines", "Product name, weight in g/kg, price per weight unit and VAT only when explicitly added."),
    ("Amounts", "Subtotal excluding VAT, VAT and total calculated automatically."),
    ("Presentation", "Company and client details, language, currency, template, terms and optional notes."),
    ("Reference and status", "Reference assigned at issue; draft, issued, sent, accepted, rejected, expired or cancelled."),
], [4.3, 13.1])
h("Creating and changing an estimate")
bullet("A draft can be completed and corrected. Issuing requires at least one valid line and the required company/client information.")
bullet("After issue, the commercial content and PDF are preserved. Corrections do not silently overwrite the original: cancel it and create a replacement.")
h("Creating invoices from an estimate")
bullet("Only an accepted estimate can be used as the source of a new invoice. Its lines are copied into an invoice draft for review before issue.")
bullet("One estimate can produce several invoices. Previous invoices, including cancelled ones, are preserved and remain linked to the estimate.")
bullet("Each invoice can have only one source estimate. Combining several estimates into one invoice is not included.")
bullet("The link identifies the invoice’s origin. It does not automatically cap the combined amount invoiced against that estimate.")
h("Please confirm")
question("E1", "Do you confirm multiple invoices per estimate, without an automatic limit on the combined invoiced amount?")
question("E2", "What estimate validity period do you use? The proposed initial default is 15 days and can be changed.")
confirm()

page("4 Invoices")
p("An invoice can be created directly, from an accepted estimate, or from eligible delivery notes. An estimate is therefore not mandatory.")
table(["Data", "Proposed rule"], [
    ("Client and source", "One client; optionally one source estimate or already-delivered delivery-note lines."),
    ("Dates", "Invoice date and a payment due date required before issue."),
    ("Lines", "Product name, positive weight, g/kg unit, price excluding VAT and explicitly selected VAT."),
    ("Amounts and currency", "Subtotal, VAT and total; MAD by default. Calculated per line, then summed."),
    ("Printed details", "Company/client identity, template, selected bank details, notes and terms when provided."),
    ("Status", "Draft, issued, sent or cancelled. Unpaid, partially paid and paid are tracked separately."),
], [4.3, 13.1])
h("Issuing and correcting invoices")
bullet("Drafts can be edited or deleted. Issuing freezes the number, identity details and printed content. An issued invoice cannot be deleted or edited like a draft.")
bullet("Corrections require a reasoned cancellation and a replacement invoice. Received payments must be resolved before cancellation; this does not automatically refund money.")
bullet("Recording a payment or sending an invoice does not replace its original PDF. Delivery-note conversion cannot bill the same delivered weight twice.")
h("VAT and public references")
p("VAT is off by default and is added only by an explicit action on each line. Line amounts are rounded to two decimal places before being summed. Available rates and worked rounding examples will be approved with your accountant.")
p("Proposed format: FAC-2026-<10 random digits>, without a visible sequential counter. DEV, BL and PAY identify the other document types. Numbers are unique, assigned once and retained after cancellation. Invoice numbering still requires your accountant’s approval.")
h("Please confirm")
question("I1", "Do you confirm VAT only when requested and an editable initial payment-term default of 15 days?")
question("I2", "Does your accountant approve the proposed reference format and required invoice wording?")
confirm()

page("5 Delivery notes")
p("A delivery note, or Bon de livraison (BL), identifies the goods and weights to deliver. It has no prices or financial totals and does not reduce inventory in the software.")
table(["Data", "Proposed rule"], [
    ("Header", "Client, delivery date, company/client details, delivery address and instructions."),
    ("Lines", "Product name/reference, g/kg unit and weight; source invoice line when delivering against an invoice."),
    ("Receipt confirmation", "Recipient name and confirmation time; a reception area is available on the printed document."),
    ("Status", "Draft, prepared, delivered, acknowledged or cancelled. Reference assigned when prepared."),
], [4.3, 13.1])
h("Delivery after invoicing")
p("One invoice can have several delivery notes for partial deliveries. Total weight allocated to non-cancelled notes cannot exceed the invoiced weight. Draft notes already reserve weight; deleting or cancelling them releases that reservation.")
h("Delivery before invoicing")
p("An independent delivery note can be created before any invoice. Once delivered or acknowledged, its lines can be invoiced. Several notes for the same client can be combined, or only part of the available delivered weight can be billed.")
p("Prices are confirmed when creating the invoice because the delivery note contains none. Invoice drafts reserve the weights selected for billing. Cancelling an invoice keeps its history but releases its billed weights for a possible replacement. The same delivered weight cannot be billed twice.")
h("Corrections and receipt evidence")
p("Prepared content is preserved. Corrections use cancellation and replacement; active billing links must be resolved before cancelling a delivery note. Receipt confirmation does not overwrite the prepared PDF. A scanned signature or separate signed PDF needs an explicit requirement.")
h("Please confirm")
question("BL1", "Do you confirm delivery before invoicing, combining notes for one client and partial billing of delivered weights?")
question("BL2", "Are recipient name and date sufficient, or must signed proof of receipt be retained?")
question("BL3", "Do you handle goods returns? Returns and inventory movements are not included in this version.")
confirm()

page("6 Payments and installments")
p("An invoice can be paid in installments using different methods. Each payment belongs to one issued invoice. The software records payments; it does not move money or perform banking transactions.")
table(["Data", "Proposed rule"], [
    ("Invoice, amount and currency", "Required. Positive amount, same currency as the invoice, no payment exceeding its remaining balance."),
    ("Payment method", "Cash, bank transfer or cheque."),
    ("Transfer reference", "Required for a bank transfer; a receiving bank account can be selected."),
    ("Cheque details", "Cheque number and issuing bank required; due date and deposit date can be recorded."),
    ("Dates", "Actual payment or cheque handover date, actual collection date, and a separate system entry timestamp."),
    ("Status", "Pending, received, rejected or cancelled; optional internal note; unique PAY reference."),
], [4.3, 13.1])
h("Calculating the remaining balance")
p("Only received payments reduce the balance. For a 1,000 DH invoice, 400 DH in cash leaves 600 DH outstanding. A pending 600 DH cheque reduces that balance only after it clears.")
bullet("A cheque remains pending when handed over or deposited. A rejected cheque does not count as collected money.")
bullet("Pending cheques do not reserve the balance. If another payment settles the invoice first, clearing the cheque requires manual resolution; an overpayment is not recorded automatically.")
bullet("Late entry is allowed: cash received on 31 January and entered on 2 February keeps 31 January as its collection date. The actual entry timestamp remains traceable.")
bullet("Received payments cannot be silently edited or deleted. Correct an error through a reasoned cancellation and a replacement entry. This does not refund the client.")
h("Please confirm")
question("PAY1", "Do you receive one payment covering several invoices, or advances before an invoice exists? Neither is included.")
question("PAY2", "Do you need future installment schedules, separate payment receipts, refunds or credit notes? These are not included.")
confirm()

page("7 Approval and outstanding details")
p("Please return this document with your section decisions and requested corrections. We will review any differences together before changing the agreed scope or implementation rules.")
h("Functional scope for approval")
p("The first version covers weighted products and categories, individual/company clients, estimates, invoices, both delivery workflows and installment tracking by cash, bank transfer or cheque.")
p("Please explicitly flag requests for inventory, returns, discounts, client-specific pricing, packaged-unit sales, advances, multi-invoice payments, future installment schedules, separate receipts, credit notes or refunds. Mentioning a request does not automatically add it to the scope.")
h("Information to provide")
bullet("An example of your current estimate, invoice and delivery note, showing the wording and fields you expect.")
bullet("Your company details, bank details to print and any VAT rates you want available.")
bullet("Your accountant’s confirmation of numbering, required wording and worked rounding examples before launch.")
h("Client decision")
p("☐ I approve the rules described without changes.")
p("☐ I approve subject to the changes listed below.")
p("☐ Further discussion is needed before approval.")
p("Affected question numbers and requested changes:")
for _ in range(3):
    p("...............................................................................................................................")
p("Name and role: .........................................................................................................")
p("Company: ................................................................................................................")
p("Date: ....................................    Signature or written approval: ................................")
p("Unresolved points remain open. This approval confirms functional rules; any extension will be discussed separately before implementation.")

doc.save(OUTPUT)
print(OUTPUT)
